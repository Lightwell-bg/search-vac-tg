"""Build data/profile.json and data/profile.md from the owner's materials (deterministic, no LLM).

Sources: everything under ``materials/`` plus paths listed in ``materials/extra_paths.txt``.
Formats: .md .txt .pdf (text layer only, no OCR) .docx. Technologies are found with the
TAXONOMY below; projects come from "## name — summary" sections and ``Portfolio_*.pdf``.
An optional LLM pass (``enrich_with_llm``) only polishes the summary/services from the
already compact profile, never from the full documents.
"""
from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from .loader import PROVISIONAL_PROFILE, compact_profile

log = logging.getLogger(__name__)

MAX_DOC_CHARS = 50_000
MIN_PDF_CHARS = 50
EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}
SKIP_DIRS = {"node_modules", "venv", "env", "__pycache__", "site-packages", "dist", "build"}
CV_NAME_RE = re.compile(r"resume|(?<![a-z])cv(?![a-z])|резюме|portfolio|портфолио", re.I)
PORTFOLIO_PDF_RE = re.compile(r"^portfolio_.*\.pdf$", re.I)

# canonical technology -> lowercase regex variants (matched on lowercased text)
TAXONOMY: dict[str, list[str]] = {
    "Python": [r"python", r"питон"],
    "aiogram": [r"aiogram"],
    "Telethon": [r"telethon"],
    "pyTelegramBotAPI": [r"pytelegrambotapi", r"telebot"],
    "Telegram Bot API": [r"telegram bot api", r"bot api", r"telegram[- ]бот", r"telegram bot", r"телеграм[- ]?бот"],
    "n8n": [r"n8n"],
    "Make": [r"make\.com", r"integromat", r"make scenario", r"сценари\w* make"],
    "Zapier": [r"zapier"],
    "OpenAI API": [r"openai", r"chatgpt", r"gpt-?[345]"],
    "Anthropic API": [r"anthropic", r"claude(?!\s*code)"],
    "Gemini API": [r"gemini"],
    "DeepSeek": [r"deepseek"],
    "OpenRouter": [r"openrouter"],
    "GigaChat": [r"gigachat", r"гигачат"],
    "LangChain": [r"langchain", r"llamaindex"],
    "RAG": [r"rag", r"retrieval[- ]augmented"],
    "pgvector": [r"pgvector"],
    "MCP": [r"mcp", r"model context protocol"],
    "AI agents": [r"ai[- ]agents?", r"llm[- ]agents?", r"ии[- ]агент\w*", r"агентн\w+ систем\w*", r"ai-агент\w*"],
    "FastAPI": [r"fastapi"],
    "Flask": [r"flask"],
    "Django": [r"django"],
    "Jinja2": [r"jinja2?"],
    "HTMX": [r"htmx"],
    "WordPress": [r"wordpress", r"вордпресс", r"wp-плагин", r"wp plugin"],
    "WooCommerce": [r"woocommerce", r"woo commerce"],
    "PHP": [r"php"],
    "JavaScript": [r"javascript", r"vanilla js"],
    "TypeScript": [r"typescript"],
    "React": [r"react"],
    "Next.js": [r"next\.js", r"nextjs"],
    "Vite": [r"vite"],
    "Tailwind": [r"tailwind"],
    "HTML/CSS": [r"html", r"css"],
    "PostgreSQL": [r"postgres\w*"],
    "SQLite": [r"sqlite", r"aiosqlite"],
    "Supabase": [r"supabase"],
    "MySQL": [r"mysql", r"mariadb"],
    "Redis": [r"redis"],
    "SQLAlchemy": [r"sqlalchemy"],
    "Alembic": [r"alembic"],
    "Docker": [r"docker"],
    "Docker Compose": [r"docker[- ]compose", r"docker-compose"],
    "Ubuntu VPS": [r"ubuntu", r"vps"],
    "Caddy": [r"caddy"],
    "Nginx": [r"nginx"],
    "systemd": [r"systemd"],
    "Google Sheets": [r"google sheets", r"gspread", r"гугл[- ]таблиц\w*", r"google таблиц\w*"],
    "Google Drive": [r"google drive", r"гугл[- ]диск"],
    "CRM (amoCRM/Bitrix24)": [r"amocrm", r"amo crm", r"bitrix24", r"битрикс ?24", r"crm"],
    "BeautifulSoup": [r"beautifulsoup", r"bs4"],
    "Playwright": [r"playwright"],
    "Selenium": [r"selenium"],
    "scraping": [r"scraping", r"scraper", r"парсинг", r"парсер\w*", r"crawler"],
    "APScheduler": [r"apscheduler"],
    "Celery": [r"celery"],
    "pytest": [r"pytest"],
    "Pydantic": [r"pydantic"],
    "ElevenLabs": [r"elevenlabs"],
    "Whisper": [r"whisper"],
    "Stripe": [r"stripe"],
    "YooKassa": [r"yookassa", r"юkassa", r"юкасса", r"yoomoney"],
    "Instagram API": [r"instagram api", r"instagram graph", r"instagrapi"],
    "VK API": [r"vk api", r"vk_api", r"vkontakte api", r"вк api", r"vk bot", r"вконтакте api"],
    "Facebook API": [r"facebook api", r"facebook graph", r"meta graph", r"facebook ads api"],
    "webhooks": [r"webhooks?", r"вебхук\w*"],
    "REST API": [r"rest api", r"restful", r"rest-api"],
}

_COMPILED: dict[str, re.Pattern] = {
    name: re.compile(r"(?<![a-zа-яё0-9_])(?:" + "|".join(v) + r")(?![a-z0-9_])")
    for name, v in TAXONOMY.items()
}

# technology triggers -> what the owner can be offered for
_RULES: list[dict] = [
    {"any": {"aiogram", "Telethon", "pyTelegramBotAPI", "Telegram Bot API"},
     "skills": ["Telegram bots"], "services": ["Telegram bot development"],
     "strong": ["Telegram bots"], "types": ["telegram_bot"]},
    {"any": {"n8n", "Make", "Zapier"},
     "skills": ["workflow automation"], "services": ["n8n / workflow automation"],
     "strong": ["workflow automation"], "types": ["automation"]},
    {"any": {"OpenAI API", "Anthropic API", "Gemini API", "DeepSeek", "OpenRouter", "GigaChat", "LangChain"},
     "skills": ["AI assistants", "LLM integrations"], "services": ["AI assistants and LLM integrations"],
     "strong": ["AI/LLM integration"], "types": ["ai_assistant"]},
    {"any": {"RAG", "pgvector"},
     "skills": ["RAG"], "services": ["RAG"], "strong": ["RAG"], "types": []},
    {"any": {"AI agents", "MCP"},
     "skills": ["AI agents", "MCP"], "services": ["AI agents"], "strong": ["AI agents"], "types": []},
    {"any": {"FastAPI", "Flask", "Django"},
     "skills": ["Python web backend"], "services": ["Python web backends"],
     "acceptable": ["Python web backend"], "types": ["web_backend"]},
    {"any": {"WordPress", "WooCommerce", "PHP"},
     "skills": ["WordPress"], "services": ["WordPress/WooCommerce customization"],
     "acceptable": ["WordPress/WooCommerce"], "types": ["wordpress"]},
    {"any": {"BeautifulSoup", "Playwright", "Selenium", "scraping"},
     "skills": ["parsing/scraping"], "services": ["parsers and scrapers"],
     "strong": ["parsing"], "types": ["parser"]},
    {"any": {"Google Sheets", "Google Drive"},
     "skills": ["Google Sheets integrations"], "services": ["Google Sheets integrations"],
     "acceptable": ["Google Sheets automation"], "types": []},
    {"any": {"CRM (amoCRM/Bitrix24)"},
     "skills": ["CRM integrations"], "services": ["CRM integrations"],
     "acceptable": ["CRM integration"], "types": []},
    {"any": {"webhooks", "REST API"},
     "skills": ["API integrations", "webhooks"], "services": ["API and webhook integrations"],
     "strong": ["API integration"], "types": []},
    {"any": {"Docker", "Docker Compose", "Ubuntu VPS", "systemd", "Caddy", "Nginx"},
     "skills": ["deployment"], "services": ["deployment on Ubuntu VPS with Docker"],
     "acceptable": ["Docker/VPS deployment"], "types": []},
    {"any": {"Stripe", "YooKassa"},
     "skills": ["payment integrations"], "services": [], "acceptable": ["payment integrations"], "types": []},
    {"any": {"Instagram API", "VK API", "Facebook API"},
     "skills": ["social media APIs"], "services": [], "acceptable": ["social media APIs"], "types": []},
    {"any": {"ElevenLabs", "Whisper"},
     "skills": ["voice AI"], "services": [], "acceptable": ["speech/voice AI integrations"], "types": []},
    {"any": {"Python"},
     "skills": ["Python"], "services": [], "strong": ["Python automation"], "types": []},
    {"any": {"React", "Next.js", "Vite", "Tailwind", "TypeScript", "JavaScript", "HTML/CSS"},
     "skills": ["web frontend as part of a project"], "services": [], "acceptable": [], "types": []},
]

REJECT_CATEGORIES = ["design", "SMM", "copywriting", "video editing", "sales", "recruiting",
                     "accounting", "non-technical tasks"]
FIXED_WEAK = ["mobile apps", "1C"]


# --------------------------------------------------------------------------- documents

@dataclass
class Document:
    path: Path
    text: str

    @property
    def name(self) -> str:
        return self.path.name


def find_technologies(text: str) -> dict[str, int]:
    """{canonical technology: number of mentions} for one text."""
    low = (text or "").lower()
    found: dict[str, int] = {}
    for name, pattern in _COMPILED.items():
        n = len(pattern.findall(low))
        if n:
            found[name] = n
    return found


MAX_PDF_PAGES = 50
MAX_DOCX_ENTRIES = 1000
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_DOCX_RATIO = 100
SUSPICIOUS_ERROR = "Файл выглядит подозрительно (слишком большой после распаковки)"


class SuspiciousFileError(ValueError):
    """Archive-like document that looks like a zip bomb."""


def check_docx_archive(path: Path) -> None:
    """Reject a DOCX (zip) with too many entries, too much data or an extreme compression ratio."""
    import zipfile
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as e:
        raise ValueError("Повреждённый DOCX") from e
    if len(infos) > MAX_DOCX_ENTRIES:
        raise SuspiciousFileError(SUSPICIOUS_ERROR)
    if sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED:
        raise SuspiciousFileError(SUSPICIOUS_ERROR)
    for i in infos:
        if i.file_size > 0 and i.file_size > MAX_DOCX_RATIO * max(i.compress_size, 1):
            raise SuspiciousFileError(SUSPICIOUS_ERROR)


def read_document(path: Path, strict: bool = False) -> str | None:
    """Text of a supported file, capped; None if unreadable or scanned (with a warning in the log).

    With ``strict`` a suspicious (zip-bomb like) DOCX raises ``SuspiciousFileError`` instead of None.
    """
    ext = path.suffix.lower()
    try:
        if ext in (".md", ".txt"):
            with open(path, encoding="utf-8-sig", errors="replace") as f:
                text = f.read(MAX_DOC_CHARS)
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            chunks: list[str] = []
            total = 0
            for n, page in enumerate(reader.pages):
                if n >= MAX_PDF_PAGES:
                    break
                chunks.append(page.extract_text() or "")
                total += len(chunks[-1])
                if total > MAX_DOC_CHARS:
                    break
            text = "\n".join(chunks)
            if len(text.strip()) < MIN_PDF_CHARS:
                log.warning("PDF has no text layer (scanned?), skipped: %s", path)
                return None
        elif ext == ".docx":
            check_docx_archive(path)
            import docx
            d = docx.Document(str(path))
            parts: list[str] = []
            total = 0
            for p in d.paragraphs:
                parts.append(p.text)
                total += len(p.text)
                if total > MAX_DOC_CHARS:
                    break
            if total <= MAX_DOC_CHARS:
                for table in d.tables:
                    for row in table.rows:
                        line = " | ".join(c.text.strip() for c in row.cells)
                        parts.append(line)
                        total += len(line)
                        if total > MAX_DOC_CHARS:
                            break
                    if total > MAX_DOC_CHARS:
                        break
            text = "\n".join(parts)
        else:
            return None
    except SuspiciousFileError:
        if strict:
            raise
        log.warning("suspicious document skipped: %s", path)
        return None
    except Exception as e:  # noqa: BLE001 - one broken file must not stop the build
        log.warning("cannot read %s: %s: %s", path, type(e).__name__, e)
        return None
    return text[:MAX_DOC_CHARS]


def _walk(root: Path) -> list[Path]:
    out: list[Path] = []
    for p in sorted(root.rglob("*")):
        rel = p.relative_to(root).parts
        if any(part.startswith(".") or part.lower() in SKIP_DIRS for part in rel):
            continue
        if p.is_file() and p.suffix.lower() in EXTENSIONS:
            out.append(p)
    return out


def collect_paths(materials_dir: Path) -> list[Path]:
    """Files to read: materials/ recursively + entries of materials/extra_paths.txt."""
    materials_dir = Path(materials_dir)
    paths: list[Path] = []
    if materials_dir.is_dir():
        for p in _walk(materials_dir):
            if p.name.lower() == "extra_paths.txt":
                continue
            if p.name == "README.md" and p.parent == materials_dir:
                continue
            paths.append(p)
    else:
        log.warning("materials dir not found: %s", materials_dir)
    extra = materials_dir / "extra_paths.txt"
    if extra.is_file():
        for line in extra.read_text(encoding="utf-8-sig", errors="replace").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            p = Path(line)
            if p.is_dir():
                paths.extend(_walk(p))
            elif p.is_file():
                if p.suffix.lower() in EXTENSIONS and not p.name.startswith("."):
                    paths.append(p)
            else:
                log.info("extra path not found, skipped: %s", line)
    seen: set[str] = set()
    unique: list[Path] = []
    for p in paths:
        key = str(p.resolve()).lower()
        if key not in seen:
            seen.add(key)
            unique.append(p)
    return unique


def collect_documents(materials_dir: Path) -> list[Document]:
    docs: list[Document] = []
    for p in collect_paths(materials_dir):
        text = read_document(p)
        if text is not None and text.strip():
            docs.append(Document(p, text))
    return docs


# ---------------------------------------------------------------------------- projects

_FIELD_RE = re.compile(r"^[\s>*_-]*(тип|стек|stack|type|вероятно чужой код)[*_\s]*:[*_\s]*(.*)$", re.I)


def split_top_level(s: str) -> list[str]:
    """Split on commas that are not inside brackets."""
    out, depth, cur = [], 0, []
    for ch in s:
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            depth = max(0, depth - 1)
        if ch == "," and depth == 0:
            out.append("".join(cur).strip())
            cur = []
        else:
            cur.append(ch)
    out.append("".join(cur).strip())
    return [x for x in out if x]


def parse_projects_md(text: str, source: str) -> list[dict]:
    """Projects from "## name — summary" sections with "Стек:"; skips "Вероятно чужой код: да"."""
    projects: list[dict] = []
    for block in re.split(r"(?m)^(?=## )", text):
        if not block.startswith("## "):
            continue
        lines = block.splitlines()
        heading = lines[0][3:].strip()
        name, _, summary = heading.partition(" — ")
        fields: dict[str, str] = {}
        for line in lines[1:]:
            m = _FIELD_RE.match(line)
            if m:
                key = m.group(1).lower()
                key = {"тип": "type", "стек": "stack", "stack": "stack", "type": "type"}.get(key, "foreign")
                fields.setdefault(key, m.group(2).strip())
        if "stack" not in fields:
            continue
        if fields.get("foreign", "").lower().startswith("да"):
            log.info("project %r skipped: probably not owner's code", name)
            continue
        stack_raw = fields["stack"]
        if not stack_raw or stack_raw.lower().startswith("нет данных"):
            continue
        projects.append({
            "name": name.strip(), "summary": summary.strip(), "type": fields.get("type", ""),
            "stack": split_top_level(stack_raw), "source": source,
        })
    return projects


_SKIP_HEADING_RE = re.compile(
    r"контакт|обо мне|о себе|навык|skills|contacts?|\babout\b|опыт|образован|оглавлен|содержан|"
    r"table of contents|\bcontents\b|резюме|education|experience|портфолио$|^portfolio$", re.I)
_EMOJI_RE = re.compile("[\U00010000-\U0010ffff☀-➿⬀-⯿️‍]")
_HEADING_NUM_RE = re.compile(r"^[\W_]*(?:\d+(?:\.\d+)*[.)]?\s+)?")
_STACK_FIELD_RE = re.compile(r"(?im)^[\s>*_-]*(?:стек|stack)[*_\s]*:")
_MD_LINK_RE = re.compile(r"!?\[([^\]]*)\]\([^)]*\)")
MIN_SECTION_CHARS = 80


def _clean_heading(h: str) -> str:
    h = _EMOJI_RE.sub("", h)
    h = _HEADING_NUM_RE.sub("", h.strip())
    return re.sub(r"[*_`]+", "", h).strip()[:80].strip()


def _plain_text(body: str) -> str:
    t = _MD_LINK_RE.sub(r"\1", body)
    t = re.sub(r"(?m)^[\s>#*_|-]+", "", t)
    t = re.sub(r"[*_`|]+", " ", t)
    return re.sub(r"\s+", " ", _EMOJI_RE.sub("", t)).strip()


def parse_sections(text: str, source: str) -> list[dict]:
    """Projects from "##"/"###" sections that have no "Стек:" line: stack = taxonomy hits in the text."""
    projects: list[dict] = []
    for block in re.split(r"(?m)^(?=#{2,3} )", text):
        m = re.match(r"#{2,3} (.*)", block)
        if not m:
            continue
        body = block[m.end():]
        if _STACK_FIELD_RE.search(body):
            continue            # handled by parse_projects_md
        name = _clean_heading(m.group(1))
        if not name or _SKIP_HEADING_RE.search(name):
            continue
        plain = _plain_text(body)
        if len(plain) < MIN_SECTION_CHARS:
            continue
        hits = find_technologies(name + " " + plain)
        if not hits:
            continue
        stack = sorted(hits, key=lambda t: (-hits[t], t))
        types = derive_from_technologies(set(stack))["types"]
        projects.append({"name": name, "summary": plain[:300], "type": types[0] if types else "",
                         "stack": stack, "source": source})
    return projects


_GENERIC_LINE_RE = re.compile(r"^[\W_]*(?:(?:kwork|portfolio|case|development|кейс|портфолио|ai)[\W_]*)+$", re.I)


def parse_portfolio_pdf(doc: Document) -> dict | None:
    lines = [ln.strip() for ln in doc.text.splitlines() if ln.strip()]
    if not lines:
        return None
    idx = 0
    # skip generic banner lines such as "KWORK PORTFOLIO CASE" when a real title follows
    while idx < len(lines) - 1 and _GENERIC_LINE_RE.match(lines[idx]):
        idx += 1
    rest = " ".join(lines[idx + 1:])
    return {
        "name": lines[idx][:80], "summary": re.sub(r"\s+", " ", rest)[:300], "type": "",
        "stack": list(find_technologies(doc.text)), "source": str(doc.path),
    }


def extract_projects(docs: list[Document]) -> list[dict]:
    projects: list[dict] = []
    seen: set[str] = set()

    def add(p: dict) -> None:
        key = re.sub(r"\W+", " ", p["name"].lower()).strip()
        if key and key not in seen:
            seen.add(key)
            projects.append(p)

    for d in docs:
        if d.path.suffix.lower() == ".pdf":
            if PORTFOLIO_PDF_RE.match(d.name):
                p = parse_portfolio_pdf(d)
                if p:
                    add(p)
        elif d.path.suffix.lower() in (".md", ".txt"):
            if d.path.suffix.lower() == ".md":
                for p in parse_projects_md(d.text, str(d.path)):
                    add(p)
            for p in parse_sections(d.text, d.name):
                add(p)
    return projects


# ---------------------------------------------------------------------------- profile

def _dedup(items: list[str]) -> list[str]:
    seen, out = set(), []
    for i in items:
        if i and i.lower() not in seen:
            seen.add(i.lower())
            out.append(i)
    return out


def derive_from_technologies(techs: set[str]) -> dict[str, list[str]]:
    """Fixed mapping technologies -> skills/services/matches/project types."""
    res: dict[str, list[str]] = {k: [] for k in ("skills", "services", "strong", "acceptable", "types")}
    for rule in _RULES:
        if rule["any"] & techs:
            for key in res:
                res[key].extend(rule.get(key, []))
    return {k: _dedup(v) for k, v in res.items()}


def build_profile(materials_dir: Path) -> dict:
    docs = collect_documents(materials_dir)
    doc_counts: dict[str, int] = {}
    mentions: dict[str, int] = {}
    for d in docs:
        for tech, n in find_technologies(d.text).items():
            doc_counts[tech] = doc_counts.get(tech, 0) + 1
            mentions[tech] = mentions.get(tech, 0) + n
    from_materials = sorted(doc_counts, key=lambda t: (-doc_counts[t], -mentions[t], t))

    base = PROVISIONAL_PROFILE
    technologies = list(from_materials)
    found_set = set(from_materials)
    weak_techs: list[str] = []
    for t in base["technologies"]:
        canon = set(find_technologies(t))
        if (canon and canon <= found_set) or t in found_set:
            continue           # already covered by materials
        technologies.append(t)
        if not (canon & found_set):
            weak_techs.append(t)

    derived = derive_from_technologies(found_set)
    projects = extract_projects(docs)
    has_cv = any(CV_NAME_RE.search(d.name) for d in docs)

    def merged(key: str, extra: list[str]) -> list[str]:
        return _dedup(extra + list(base.get(key, [])))

    profile = {
        "provisional": not (has_cv or projects),
        "summary": "",
        "skills": merged("skills", derived["skills"]),
        "technologies": _dedup(technologies),
        "technology_documents": {t: doc_counts[t] for t in from_materials},
        "services": merged("services", derived["services"]),
        "project_types": merged("project_types", derived["types"]),
        "portfolio_projects": projects,
        "strong_matches": merged("strong_matches", derived["strong"]),
        "acceptable_matches": merged("acceptable_matches", derived["acceptable"]),
        "weak_matches": _dedup(weak_techs + FIXED_WEAK),
        "reject_categories": list(REJECT_CATEGORIES),
        "sources": [str(d.path) for d in docs],
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }
    top = from_materials[:8] or base["technologies"][:8]
    profile["summary"] = (f"Developer ({', '.join(profile['strong_matches'][:4])}). "
                          f"Main stack: {', '.join(top)}.")
    return profile


# ------------------------------------------------------------------------------- LLM

async def enrich_with_llm(profile: dict, llm, repo=None) -> dict:
    """Ask OpenRouter for a 2-sentence Russian summary and up to 8 services.

    Only the compact profile and project names/stacks are sent, never the full documents.
    """
    from ..llm.schemas import LlmError

    projects = "\n".join(f"- {p['name']}: {', '.join(p.get('stack', [])[:8])}"
                         for p in profile.get("portfolio_projects", [])[:40])
    system = ("You help build a freelancer profile. Reply with a JSON object only: "
              '{"summary": "<2 sentences in Russian>", "services": ["up to 8 short service names"]}. '
              "Use only the given data; do not invent facts.")
    user = f"Compact profile:\n{compact_profile(profile)}\n\nProjects:\n{projects}"
    try:
        content, usage = await llm.chat(system, user, json_mode=True, max_tokens=500, temperature=0.2)
    except LlmError as e:
        u = e.usage
        if repo is not None:
            await repo.log_llm_usage(purpose="profile_building", model=(u.model if u else llm.model) or "",
                                     input_tokens=u.input_tokens if u else 0,
                                     output_tokens=u.output_tokens if u else 0,
                                     cost_usd=u.cost_usd if u else 0.0,
                                     duration_ms=u.duration_ms if u else 0, error=str(e)[:500])
        log.warning("LLM profile enrichment failed: %s", e)
        return profile
    if repo is not None:
        await repo.log_llm_usage(purpose="profile_building", model=usage.model,
                                 input_tokens=usage.input_tokens, output_tokens=usage.output_tokens,
                                 cost_usd=usage.cost_usd, duration_ms=usage.duration_ms)
    from ..llm.openrouter import parse_json_object
    try:
        data = parse_json_object(content)
    except LlmError as e:
        log.warning("LLM profile enrichment returned bad JSON: %s", e)
        return profile
    if isinstance(data.get("summary"), str) and data["summary"].strip():
        profile["summary"] = data["summary"].strip()
    services = data.get("services")
    if isinstance(services, list):
        profile["services"] = _dedup([str(s).strip() for s in services[:8]] + profile["services"])
    return profile


# ------------------------------------------------------------------------------ output

def render_markdown(profile: dict) -> str:
    def bullets(items) -> list[str]:
        return [f"- {i}" for i in items] or ["- (нет данных)"]

    out = ["# Профиль исполнителя", "", profile.get("summary", ""), ""]
    if profile.get("provisional"):
        out += ["> Профиль предварительный: резюме/портфолио не найдены. Положите файлы в `materials/` "
                "и запустите `python scripts/rebuild_profile.py`.", ""]
    out += [f"Сгенерирован: {profile.get('generated_at', '')}", ""]
    for title, key in (("Технологии", "technologies"), ("Услуги", "services"), ("Навыки", "skills"),
                       ("Сильное совпадение", "strong_matches"), ("Приемлемое совпадение", "acceptable_matches"),
                       ("Слабое совпадение", "weak_matches"), ("Не берём", "reject_categories")):
        out += [f"## {title}", *bullets(profile.get(key, [])), ""]
    out += ["## Проекты портфолио"]
    projects = profile.get("portfolio_projects", [])
    if not projects:
        out.append("- (нет данных)")
    for p in projects:
        line = f"- **{p['name']}**"
        if p.get("summary"):
            line += f" — {p['summary']}"
        if p.get("stack"):
            line += f" (стек: {', '.join(p['stack'][:10])})"
        out.append(line)
    out += ["", "## Источники", *bullets(profile.get("sources", [])), ""]
    return "\n".join(out)


def write_profile(profile: dict, json_path: Path, md_path: Path) -> None:
    for p in (json_path, md_path):
        Path(p).parent.mkdir(parents=True, exist_ok=True)
    Path(json_path).write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    Path(md_path).write_text(render_markdown(profile), encoding="utf-8")
