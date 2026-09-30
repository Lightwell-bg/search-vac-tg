import json

import docx

from src.profile.builder import (
    build_profile,
    find_technologies,
    parse_projects_md,
    split_top_level,
    write_profile,
)

PROJECTS_MD = """# Проекты

## botshop — Магазин в Telegram
- Тип: telegram_bot
- Стек: Python, aiogram 3, PostgreSQL + pgvector, WordPress-плагин (PHP, vanilla JS)
- Вероятно чужой код: нет

## foreign-tool — Чужая утилита
- Тип: automation
- Стек: Python, Selenium
- Вероятно чужой код: да (папка -master)

## empty — пустая папка
- Тип: неизвестно
- Стек: нет данных
- Вероятно чужой код: нет

## Найденные файлы резюме
Просто список без стека.
"""


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_find_technologies_boundaries():
    found = find_technologies("Built a Telegram bot on aiogram with FastAPI, n8n and Docker Compose. Rage of Reactive.")
    assert {"aiogram", "FastAPI", "n8n", "Docker Compose", "Telegram Bot API"} <= set(found)
    assert "RAG" not in found and "React" not in found
    assert "Anthropic API" not in find_technologies("we use Claude Code daily")


def test_split_top_level_keeps_brackets():
    assert split_top_level("Python, WordPress-плагин (PHP, vanilla JS), Docker") == [
        "Python", "WordPress-плагин (PHP, vanilla JS)", "Docker"]


def test_projects_md_skips_foreign_empty_and_stackless():
    projects = parse_projects_md(PROJECTS_MD, "x.md")
    assert [p["name"] for p in projects] == ["botshop"]
    p = projects[0]
    assert p["summary"] == "Магазин в Telegram" and p["type"] == "telegram_bot"
    assert p["stack"] == ["Python", "aiogram 3", "PostgreSQL + pgvector", "WordPress-плагин (PHP, vanilla JS)"]


def test_build_profile_from_mixed_materials(tmp_path):
    m = tmp_path / "materials"
    (m / "sub").mkdir(parents=True)
    (m / "README.md").write_text("aiogram Telethon n8n", encoding="utf-8")   # ignored: materials README
    (m / ".hidden.md").write_text("Zapier", encoding="utf-8")                # ignored: hidden
    (m / "projects.md").write_text(PROJECTS_MD, encoding="utf-8")
    (m / "notes.txt").write_text("I build parsers with Playwright and BeautifulSoup.", encoding="utf-8")
    d = docx.Document()
    d.add_paragraph("Resume: FastAPI developer")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Redis"
    t.cell(0, 1).text = "Celery"
    d.save(m / "sub" / "resume.docx")
    (m / "Portfolio_Demo.pdf").write_bytes(
        minimal_pdf("Demo Bot Case. A Telegram bot with aiogram and PostgreSQL for a shop, "
                    "deployed with Docker on a VPS. More text to pass the length threshold."))
    extra = tmp_path / "outside.md"
    extra.write_text("Stripe payments", encoding="utf-8")
    (m / "extra_paths.txt").write_text(f"# c\n{extra}\n{tmp_path / 'missing.md'}\n", encoding="utf-8")

    prof = build_profile(m)
    techs = prof["technologies"]
    assert {"Playwright", "BeautifulSoup", "FastAPI", "Redis", "Celery", "Stripe", "aiogram"} <= set(techs)
    docs = prof["technology_documents"]
    assert "Zapier" not in techs and "Telethon" not in docs and "n8n" not in docs and "n8n" in techs
    assert techs.index("aiogram") < techs.index("Gemini API")                   # materials first
    names = [p["name"] for p in prof["portfolio_projects"]]
    assert "botshop" in names and "foreign-tool" not in names
    assert any(n.startswith("Demo Bot Case") for n in names)
    assert prof["provisional"] is False
    assert not any(s.endswith("README.md") or ".hidden" in s for s in prof["sources"])
    assert "Telegram bot development" in prof["services"]
    assert "n8n" in prof["weak_matches"] and "1C" in prof["weak_matches"]
    assert prof["reject_categories"][0] == "design"


def test_provisional_flag_rules(tmp_path):
    empty = tmp_path / "m1"
    empty.mkdir()
    assert build_profile(empty)["provisional"] is True

    m2 = tmp_path / "m2"
    m2.mkdir()
    (m2 / "notes.md").write_text("python and docker", encoding="utf-8")
    assert build_profile(m2)["provisional"] is True            # no cv-like name, no projects

    m3 = tmp_path / "m3"
    m3.mkdir()
    (m3 / "CV_ivan.md").write_text("python", encoding="utf-8")
    assert build_profile(m3)["provisional"] is False           # cv-like file name

    m4 = tmp_path / "m4"
    m4.mkdir()
    (m4 / "cases.md").write_text(PROJECTS_MD, encoding="utf-8")
    assert build_profile(m4)["provisional"] is False           # projects extracted


def test_write_profile_files(tmp_path):
    m = tmp_path / "m"
    m.mkdir()
    (m / "resume.md").write_text("Python, aiogram", encoding="utf-8")
    prof = build_profile(m)
    write_profile(prof, tmp_path / "out" / "profile.json", tmp_path / "out" / "profile.md")
    loaded = json.loads((tmp_path / "out" / "profile.json").read_text(encoding="utf-8"))
    assert loaded["technologies"][0] in ("Python", "aiogram")
    assert "Технологии" in (tmp_path / "out" / "profile.md").read_text(encoding="utf-8")
